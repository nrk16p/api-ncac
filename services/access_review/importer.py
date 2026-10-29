"""Access Review — import ATMS role groups / Winspeed security / Email groups from Excel.

parse_atms / parse_winspeed / parse_email are pure (take an openpyxl Workbook, read_only or not).
run_import resolves employees against `users`, diffs against the ar_* tables and — only when dry_run is False —
writes: upsert roles/user_roles, sync the items of every role present in the file, deactivate
active user_roles of that system missing from the file.

Item sync is a diff, so item ids stay stable (review drafts keep pointing at the same items): new items are
inserted, items back in the file are reactivated, and an item gone from the file is deleted — unless a review
references it (ar_review_line_items), then it is set is_active=false (hidden from new reviews, kept for history). Dry run issues SELECTs only.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Dict, List, Optional, Set, Tuple

import openpyxl
from sqlalchemy import insert
from sqlalchemy.orm import Session

from models.access_review_model import ArRole, ArRoleItem, ArSystem, ArUserRole
from services.access_review import repo

SYSTEMS = ("ATMS", "WINSPEED", "EMAIL")
SYSTEM_NAMES = {"ATMS": "ATMS", "WINSPEED": "Winspeed", "EMAIL": "Email Group"}

ATMS_INDEX_SHEETS = {"สรุป"}
ATMS_HEADER_ROW = 5
WAREHOUSE_CATEGORY = "คลังสินค้า"

REASON_NO_USERNAME = "ไม่พบชื่อผู้ใช้นี้ในระบบ (users.username / email)"
REASON_NO_EMPLOYEE_ID = "ผู้ใช้ไม่มีรหัสพนักงาน"
REASON_UNKNOWN_EMPLOYEE = "ไม่พบรหัสพนักงานในระบบ"
REASON_MISSING_EMPLOYEE = "ไม่มีรหัสพนักงานในไฟล์"
REASON_MISSING_SHEET = "ไม่พบชีทของผู้ใช้"
USER_GRANT_PREFIX = "สิทธิ์เฉพาะบุคคล"

Item = Tuple[str, Optional[str], Optional[str]]  # (category, item, permission)


class ImportFileError(ValueError):
    """The workbook does not look like the expected export."""


@dataclass
class RawAssignment:
    role_name: str
    account_name: Optional[str] = None
    employee_id: Optional[str] = None  # WINSPEED / EMAIL
    username: Optional[str] = None     # ATMS: matched to users.username / email local-part


@dataclass
class Parsed:
    system: str
    roles: Dict[str, Optional[List[Item]]] = field(default_factory=dict)  # None → no items (EMAIL)
    assignments: List[RawAssignment] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    unmatched: List[dict] = field(default_factory=list)


# ---------------------------- cell helpers ----------------------------

def _text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _employee_id(value) -> Optional[str]:
    text = _text(value)
    if text and text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def _cell(row: tuple, idx: int):
    return row[idx] if idx < len(row) else None


def _dedupe(items: List[Item]) -> List[Item]:
    return list(dict.fromkeys(items))


def load_workbook(data: bytes):
    try:
        return openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types for a non-xlsx upload
        raise ImportFileError(f"อ่านไฟล์ Excel ไม่ได้: {exc}") from exc


# ---------------------------- parsers ----------------------------

def parse_atms(wb) -> Parsed:
    """Every sheet except 'สรุป' is one role group. Role = A1 (titles are cut at 31 chars). Row 5 = headers;
    from row 6: A = usernames until the first blank, C = warehouse, E = module, F = permission ('-' = none)."""
    parsed = Parsed("ATMS")
    for ws in wb.worksheets:
        if ws.title.strip() in ATMS_INDEX_SHEETS:
            continue
        role_name: Optional[str] = None
        items: List[Item] = []
        usernames: List[str] = []
        users_open = True
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx == 1:
                role_name = _text(_cell(row, 0))
                continue
            if idx <= ATMS_HEADER_ROW:
                continue
            username = _text(_cell(row, 0))
            if users_open:
                if username and username.startswith("(") and username.endswith(")"):
                    users_open = False  # placeholder such as "(ไม่มีผู้ใช้งาน)"
                elif username:
                    usernames.append(username)
                else:
                    users_open = False
            warehouse = _text(_cell(row, 2))
            if warehouse:
                items.append((WAREHOUSE_CATEGORY, warehouse, None))
            module = _text(_cell(row, 4))
            if module:
                perm = _text(_cell(row, 5))
                items.append((module, None, None if perm in (None, "-") else perm))
        if not role_name:
            parsed.warnings.append(f"ชีท '{ws.title}' ไม่มีชื่อกลุ่มที่ A1 — ข้าม")
            continue
        parsed.roles[role_name] = _dedupe(parsed.roles.get(role_name, []) + items)
        for u in dict.fromkeys(usernames):
            parsed.assignments.append(RawAssignment(role_name=role_name, account_name=u, username=u))
    if not parsed.roles:
        raise ImportFileError("ไม่พบชีทกลุ่มสิทธิ์ ATMS ในไฟล์")
    return parsed


def parse_winspeed(wb) -> Parsed:
    """'Summary' = [รหัสพนักงาน, Sheet, User ID, Name]. Each user sheet: header at row 4
    [Application, Object, Control, Description, Object Type, Status, Group/User, User Type, Priority].
    Only Status == 'Enabled'. Role = Group/User (User Type = User → "สิทธิ์เฉพาะบุคคล (<UserID>)");
    items category=Description, item=Control."""
    if "Summary" not in wb.sheetnames:
        raise ImportFileError("ไม่พบชีท 'Summary' ในไฟล์ Winspeed")
    parsed = Parsed("WINSPEED")
    users = []
    for idx, row in enumerate(wb["Summary"].iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        emp, sheet, user_id = _employee_id(_cell(row, 0)), _text(_cell(row, 1)), _text(_cell(row, 2))
        if sheet or user_id:
            users.append((emp, sheet, user_id))

    role_items: Dict[str, List[Item]] = {}
    for emp, sheet, user_id in users:
        if not sheet or sheet not in wb.sheetnames:
            parsed.unmatched.append({"account_name": user_id, "employee_id": emp, "role_name": None,
                                     "reason": REASON_MISSING_SHEET})
            continue
        header: Optional[Dict[str, int]] = None
        user_roles: List[str] = []
        for row in wb[sheet].iter_rows(values_only=True):
            if header is None:
                labels = [_text(c) for c in row]
                if "Status" in labels and "Group/User" in labels:
                    header = {label: i for i, label in enumerate(labels) if label}
                continue
            if _text(_cell(row, header["Status"])) != "Enabled":
                continue
            role = _text(_cell(row, header["Group/User"]))
            if not role:
                continue
            # a user-level grant (User Type = User) is a real extra permission — keep it, named so it reads as one
            if "User Type" in header and _text(_cell(row, header["User Type"])) == "User":
                role = f"{USER_GRANT_PREFIX} ({role})"
            category = _text(_cell(row, header["Description"])) if "Description" in header else None
            control = _text(_cell(row, header["Control"])) if "Control" in header else None
            role_items.setdefault(role, []).append((category or "-", control, None))
            user_roles.append(role)
        if header is None:
            parsed.warnings.append(f"ชีท '{sheet}' ไม่มีหัวตาราง Status / Group/User — ข้าม")
            continue
        for role in dict.fromkeys(user_roles):
            parsed.assignments.append(RawAssignment(role_name=role, account_name=user_id, employee_id=emp))
    parsed.roles = {name: _dedupe(items) for name, items in role_items.items()}
    return parsed


def parse_email(wb) -> Parsed:
    """First sheet: A = employee_id, B..F = group mail addresses ('-' / blank skipped). Header row skipped."""
    parsed = Parsed("EMAIL")
    ws = wb.worksheets[0]
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        emp = _employee_id(_cell(row, 0))
        if idx == 1 and (emp is None or not emp.isdigit()):
            continue
        addresses = [_text(_cell(row, i)) for i in range(1, 6)]
        addresses = [a.lower() for a in addresses if a and a != "-"]
        if not addresses:
            continue
        for addr in dict.fromkeys(addresses):
            parsed.roles.setdefault(addr, None)
            parsed.assignments.append(RawAssignment(role_name=addr, account_name=None, employee_id=emp))
    return parsed


PARSERS = {"ATMS": parse_atms, "WINSPEED": parse_winspeed, "EMAIL": parse_email}


def parse(system: str, wb) -> Parsed:
    if system not in PARSERS:
        raise ImportFileError(f"ระบบไม่ถูกต้อง: {system} (ATMS | WINSPEED | EMAIL)")
    return PARSERS[system](wb)


# ---------------------------- employee matching ----------------------------

@dataclass
class UserIndex:
    employee_ids: Set[str]
    by_login: Dict[str, Optional[str]]  # lowercased username / email local-part → employee_id

    @classmethod
    def build(cls, rows) -> "UserIndex":
        """rows: (employee_id, username, email, employee_status). Exact username beats email local-part;
        an Active user beats an inactive one."""
        ranked: Dict[str, Tuple[int, Optional[str]]] = {}

        def offer(key: Optional[str], score: int, eid: Optional[str]) -> None:
            if not key:
                return
            key = key.strip().lower()
            if key and (key not in ranked or score > ranked[key][0]):
                ranked[key] = (score, eid)

        ids: Set[str] = set()
        for eid, username, email, status in rows:
            if eid:
                ids.add(eid)
            active = 1 if (status or "").lower() == "active" else 0
            offer(username, 10 + active, eid)
            if email and "@" in email:
                offer(email.split("@", 1)[0], 2 + active, eid)
        return cls(employee_ids=ids, by_login={k: v for k, (_, v) in ranked.items()})


def resolve(parsed: Parsed, index: UserIndex) -> Tuple[List[Tuple[str, Optional[str], str]], List[dict]]:
    """→ ([(employee_id, account_name, role_name)], unmatched)."""
    matched: List[Tuple[str, Optional[str], str]] = []
    unmatched = list(parsed.unmatched)
    for a in parsed.assignments:
        reason = None
        eid = a.employee_id
        if a.username is not None:
            key = a.username.lower()
            if key not in index.by_login:
                reason = REASON_NO_USERNAME
            else:
                eid = index.by_login[key]
                if not eid:
                    reason = REASON_NO_EMPLOYEE_ID
        elif not eid:
            reason = REASON_MISSING_EMPLOYEE
        elif eid not in index.employee_ids:
            reason = REASON_UNKNOWN_EMPLOYEE
        if reason:
            unmatched.append({"account_name": a.account_name, "employee_id": eid, "role_name": a.role_name,
                              "reason": reason})
        else:
            matched.append((eid, a.account_name, a.role_name))
    return matched, unmatched


# ---------------------------- diff + write ----------------------------

def run_import(db: Session, system: str, parsed: Parsed, dry_run: bool = True) -> dict:
    index = UserIndex.build(repo.all_users_for_matching(db))
    matched, unmatched = resolve(parsed, index)

    sys_row = repo.get_system_by_code(db, system) if repo.ar_tables_exist(db) else None
    roles: Dict[str, ArRole] = {}
    existing: Dict[Tuple[str, str], ArUserRole] = {}  # (employee_id, role name) → row
    old_items: Dict[int, Dict[Item, ArRoleItem]] = {}
    if sys_row is not None:
        roles = {r.name: r for r in db.query(ArRole).filter(ArRole.system_id == sys_row.id).all()}
        role_names = {r.id: r.name for r in roles.values()}
        for ur in (db.query(ArUserRole).join(ArRole, ArRole.id == ArUserRole.role_id)
                   .filter(ArRole.system_id == sys_row.id).all()):
            existing[(ur.employee_id, role_names[ur.role_id])] = ur
        file_role_ids = [roles[n].id for n in parsed.roles if n in roles]
        if file_role_ids:
            for it in db.query(ArRoleItem).filter(ArRoleItem.role_id.in_(file_role_ids)).all():
                old_items.setdefault(it.role_id, {})[(it.category, it.item, it.permission)] = it

    wanted: Dict[Tuple[str, str], Optional[str]] = {}  # (employee_id, role name) → account_name (first wins)
    for eid, account_name, role_name in matched:
        wanted.setdefault((eid, role_name), account_name)

    report = {"system": system, "dry_run": dry_run, "roles_added": 0, "roles_updated": 0, "items_replaced": 0,
              "user_roles_in_file": len(wanted), "user_roles_added": 0, "user_roles_updated": 0,
              "user_roles_reactivated": 0, "user_roles_deactivated": 0, "unmatched": unmatched,
              # additive
              "roles_in_file": len(parsed.roles), "employees_in_file": len({eid for eid, _ in wanted}),
              "items_added": 0, "items_reactivated": 0, "items_deleted": 0, "items_kept_for_history": 0,
              "warnings": parsed.warnings}

    if sys_row is None and not dry_run:
        sys_row = ArSystem(code=system, name=SYSTEM_NAMES.get(system, system), is_active=True)
        db.add(sys_row)
        db.flush()

    # roles + items
    items_to_write: Dict[str, List[Item]] = {}
    for name, items in parsed.roles.items():
        role = roles.get(name)
        old_active = {k for k, it in old_items.get(role.id, {}).items() if it.is_active} if role else set()
        changed_items = items is not None and (role is None or set(items) != old_active)
        if role is None:
            report["roles_added"] += 1
            if not dry_run:
                role = ArRole(system_id=sys_row.id, name=name, is_active=True)
                db.add(role)
                roles[name] = role
        elif changed_items or not role.is_active:
            report["roles_updated"] += 1
            if not dry_run:
                role.is_active = True
        if changed_items:
            items_to_write[name] = items
            report["items_replaced"] += len(items)
    if not dry_run:
        db.flush()

    # items: diff against what is stored (see module docstring)
    to_insert: List[dict] = []
    to_drop: List[ArRoleItem] = []
    for name, items in items_to_write.items():
        role = roles.get(name)
        old = old_items.get(role.id, {}) if role is not None and role.id else {}
        new_keys = set(items)
        for key in items:
            it = old.get(key)
            if it is None:
                report["items_added"] += 1
                if not dry_run:
                    to_insert.append({"role_id": role.id, "category": key[0], "item": key[1], "permission": key[2],
                                      "is_active": True})
            elif not it.is_active:
                report["items_reactivated"] += 1
                if not dry_run:
                    it.is_active = True
        to_drop.extend(it for key, it in old.items() if key not in new_keys and it.is_active)
    referenced = repo.referenced_item_ids(db, [it.id for it in to_drop]) if to_drop else set()
    for it in to_drop:
        if it.id in referenced:
            report["items_kept_for_history"] += 1
            if not dry_run:
                it.is_active = False
        else:
            report["items_deleted"] += 1
            if not dry_run:
                db.delete(it)
    if not dry_run and to_insert:
        db.execute(insert(ArRoleItem), to_insert)

    # user roles
    for (eid, role_name), account_name in wanted.items():
        ur = existing.get((eid, role_name))
        if ur is None:
            report["user_roles_added"] += 1
            if not dry_run:
                db.add(ArUserRole(employee_id=eid, role_id=roles[role_name].id, account_name=account_name,
                                  is_active=True))
            continue
        if not ur.is_active:
            report["user_roles_reactivated"] += 1
        elif ur.account_name != account_name:
            report["user_roles_updated"] += 1
        if not dry_run:
            ur.is_active = True
            ur.account_name = account_name
    for key, ur in existing.items():
        if ur.is_active and key not in wanted:
            report["user_roles_deactivated"] += 1
            if not dry_run:
                ur.is_active = False

    if dry_run:
        db.rollback()
    else:
        db.commit()
    return report
