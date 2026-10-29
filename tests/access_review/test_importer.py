"""Importer — parsers on small in-memory workbooks, matching, and dry-run / commit diff on SQLite."""
from io import BytesIO

import openpyxl
import pytest

from models.access_review_model import ArRole, ArRoleItem, ArUserRole
from services.access_review import importer as I


def _roundtrip(wb):
    """Save + reopen read_only, the way the endpoint / CLI read files."""
    buf = BytesIO()
    wb.save(buf)
    return I.load_workbook(buf.getvalue())


def atms_workbook():
    wb = openpyxl.Workbook()
    wb.active.title = "สรุป"
    wb.active.append(["สรุปกลุ่มผู้ใช้งาน MENA ATMS"])
    long_name = "Chief Commercial Officer (Acting COO)"
    for title, role, users, rows in [
        ("ACC-Officer", "ACC-Officer", ["r.user", "OP"],
         [("คลัง IT", "อุบัติเหตุ", "บันทึก"), ("คลัง OPS", "อุบัติเหตุ", "ลบ"), (None, "บัญชี / การเงิน", "-"),
          (None, "อุบัติเหตุ", "บันทึก")]),
        (long_name[:31], long_name, ["Exec.Mail"], [("คลัง IT", "รายงาน", "แสดง")]),
        ("Empty", "Empty-Group", ["(ไม่มีผู้ใช้งาน)"], [("คลัง IT", None, None)]),
    ]:
        ws = wb.create_sheet(title)
        ws.append([role])
        ws.append(["ID", 1])
        ws.append(["ลิงก์"])
        ws.append([])
        ws.append([f"ผู้ใช้งาน ({len(users)})", None, "คลังสินค้า", None, "เมนู / โมดูล", "สิทธิ์"])
        n = max(len(users), len(rows))
        for i in range(n):
            u = users[i] if i < len(users) else None
            wh, mod, perm = rows[i] if i < len(rows) else (None, None, None)
            ws.append([u, None, wh, None, mod, perm])
        ws.append([None, None, None, None, None, None])
        ws.append(["late.user", None, None, None, None, None])  # after the first blank in col A → ignored
    return _roundtrip(wb)


def winspeed_workbook(extra_rows=()):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Summary"
    s.append(["รหัสพนักงาน", "Sheet", "User ID", "Name"])
    s.append([590373, "Athitiya Mayang", "Fin_LKB4", "Athitiya Mayang"])
    s.append([660548.0, "Chayada K", "Fin_BKK12", "Chayada K"])
    s.append([999999, "Missing Sheet", "Fin_X", "Nobody"])
    header = ["Application", "Object", "Control", "Description", "Object Type", "Status", "Group/User",
              "User Type", "Priority"]
    for title, uid, rows in [
        ("Athitiya Mayang", "Fin_LKB4", [("A1", "AC Accounts Payable", "Enabled", "AP ACC"),
                                         ("A2", "AC Accounts Payable", "Enabled", "AP ACC"),
                                         ("A3", "AC Accounts Payable", "Disabled", "AP ACC"),
                                         ("B1", "GL", "Enabled", "GL ACC")]),
        ("Chayada K", "Fin_BKK12", [("A1", "AC Accounts Payable", "Enabled", "AP ACC"),
                                    ("A9", "AC Accounts Payable", "Enabled", "AP ACC")]),
    ]:
        ws = wb.create_sheet(title)
        ws.append(["User ID", uid])
        ws.append(["Name", title])
        ws.append([])
        ws.append(header)
        for control, desc, status, group in rows:
            ws.append(["emapp", "Menu", control, desc, "menuitem", status, group, "Group", 1])
    for title, row in extra_rows:
        wb[title].append(row)
    return _roundtrip(wb)


def email_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["รหัส\nพนักงาน", "Group Mail1", "Group Mail2", "Group Mail3", "Group Mail4", "Group Mail5"])
    ws.append([360002, "bod@menatransport.co.th", "C-Level@menatransport.co.th", "-", "-", None])
    ws.append([493005, "bod@menatransport.co.th", "-", "-", "-", "-"])
    ws.append([111111, "-", "-", "-", "-", "-"])
    ws.append([690656, "mm@menatransport.co.th", None, None, None, None])
    return _roundtrip(wb)


# ---------------------------- parsers ----------------------------

def test_parse_atms():
    p = I.parse_atms(atms_workbook())
    assert set(p.roles) == {"ACC-Officer", "Chief Commercial Officer (Acting COO)", "Empty-Group"}  # A1, not title
    assert p.roles["ACC-Officer"] == [
        ("คลังสินค้า", "คลัง IT", None), ("อุบัติเหตุ", None, "บันทึก"), ("คลังสินค้า", "คลัง OPS", None),
        ("อุบัติเหตุ", None, "ลบ"), ("บัญชี / การเงิน", None, None)]  # '-' → None, deduped
    assert [(a.role_name, a.username) for a in p.assignments] == [
        ("ACC-Officer", "r.user"), ("ACC-Officer", "OP"), ("Chief Commercial Officer (Acting COO)", "Exec.Mail")]


def test_parse_winspeed():
    p = I.parse_winspeed(winspeed_workbook())
    assert p.roles == {
        "AP ACC": [("AC Accounts Payable", "A1", None), ("AC Accounts Payable", "A2", None),
                   ("AC Accounts Payable", "A9", None)],  # union across users, Disabled skipped
        "GL ACC": [("GL", "B1", None)],
    }
    assert sorted((a.employee_id, a.account_name, a.role_name) for a in p.assignments) == [
        ("590373", "Fin_LKB4", "AP ACC"), ("590373", "Fin_LKB4", "GL ACC"), ("660548", "Fin_BKK12", "AP ACC")]
    assert p.unmatched == [{"account_name": "Fin_X", "employee_id": "999999", "role_name": None,
                            "reason": I.REASON_MISSING_SHEET}]


def test_parse_winspeed_user_level_grant_is_its_own_role():
    p = I.parse_winspeed(winspeed_workbook(extra_rows=[
        ("Chayada K", ["emapp", "Menu", "X1", "CM Cash", "menuitem", "Enabled", "Fin_BKK12", "User", 1])]))
    assert p.roles["สิทธิ์เฉพาะบุคคล (Fin_BKK12)"] == [("CM Cash", "X1", None)]
    assert ("660548", "Fin_BKK12", "สิทธิ์เฉพาะบุคคล (Fin_BKK12)") in {
        (a.employee_id, a.account_name, a.role_name) for a in p.assignments}


def test_parse_email():
    p = I.parse_email(email_workbook())
    assert p.roles == {"bod@menatransport.co.th": None, "c-level@menatransport.co.th": None,
                       "mm@menatransport.co.th": None}
    assert [(a.employee_id, a.role_name, a.account_name) for a in p.assignments] == [
        ("360002", "bod@menatransport.co.th", None), ("360002", "c-level@menatransport.co.th", None),
        ("493005", "bod@menatransport.co.th", None), ("690656", "mm@menatransport.co.th", None)]


def test_parse_rejects_wrong_file():
    with pytest.raises(I.ImportFileError):
        I.parse_winspeed(email_workbook())
    with pytest.raises(I.ImportFileError):
        I.load_workbook(b"not an xlsx")


def test_user_index_matching():
    idx = I.UserIndex.build([
        ("1", "r.user", "r.user@menatransport.co.th", "Active"),
        ("2", "someone", "exec.mail@menatransport.co.th", "Active"),   # matched by email local-part
        ("3", "old", "exec.mail@other.co.th", "Inactive"),            # loses to the active one
    ])
    p = I.parse_atms(atms_workbook())
    matched, unmatched = I.resolve(p, idx)
    assert sorted(matched) == [("1", "r.user", "ACC-Officer"), ("2", "Exec.Mail", "Chief Commercial Officer (Acting COO)")]
    assert unmatched == [{"account_name": "OP", "employee_id": None, "role_name": "ACC-Officer",
                          "reason": I.REASON_NO_USERNAME}]


# ---------------------------- run_import (SQLite) ----------------------------

def _counts(db):
    return (db.query(ArRole).count(), db.query(ArRoleItem).count(),
            db.query(ArUserRole).filter(ArUserRole.is_active.is_(True)).count())


def test_run_import_dry_run_writes_nothing_then_commit(db, org):
    org.user("360002", 8, 1)
    org.user("493005", 3, 1)
    before = _counts(db)
    p = I.parse_email(email_workbook())
    report = I.run_import(db, "EMAIL", p, dry_run=True)
    assert _counts(db) == before
    assert (report["roles_added"], report["user_roles_in_file"], report["user_roles_added"]) == (3, 3, 3)
    assert [(u["employee_id"], u["reason"]) for u in report["unmatched"]] == [("690656", I.REASON_UNKNOWN_EMPLOYEE)]

    report = I.run_import(db, "EMAIL", p, dry_run=False)
    assert report["dry_run"] is False
    assert _counts(db) == (before[0] + 3, before[1], before[2] + 3)

    # re-import is a no-op
    again = I.run_import(db, "EMAIL", p, dry_run=True)
    assert (again["roles_added"], again["user_roles_added"], again["user_roles_updated"],
            again["user_roles_reactivated"], again["user_roles_deactivated"]) == (0, 0, 0, 0, 0)


def test_run_import_replaces_items_and_deactivates_missing(db, org):
    org.user("590373", 3, 1)
    org.user("660548", 3, 1)
    I.run_import(db, "WINSPEED", I.parse_winspeed(winspeed_workbook()), dry_run=False)
    ap = db.query(ArRole).filter(ArRole.name == "AP ACC").one()
    assert db.query(ArRoleItem).filter(ArRoleItem.role_id == ap.id).count() == 3

    # second file: AP ACC has only A1 now and Chayada lost it
    p = I.parse_winspeed(winspeed_workbook())
    p.roles["AP ACC"] = [("AC Accounts Payable", "A1", None)]
    p.assignments = [a for a in p.assignments if a.employee_id != "660548"]
    dry = I.run_import(db, "WINSPEED", p, dry_run=True)
    assert (dry["roles_updated"], dry["items_replaced"], dry["user_roles_deactivated"]) == (1, 1, 1)
    assert db.query(ArRoleItem).filter(ArRoleItem.role_id == ap.id).count() == 3  # untouched by dry run

    I.run_import(db, "WINSPEED", p, dry_run=False)
    db.expire_all()
    assert [i.item for i in db.query(ArRoleItem).filter(ArRoleItem.role_id == ap.id)] == ["A1"]
    assert [u.is_active for u in db.query(ArUserRole).filter(ArUserRole.employee_id == "660548")] == [False]

    # Chayada back with a new login → reactivated; Athitiya's login renamed → updated
    p2 = I.parse_winspeed(winspeed_workbook())
    for a in p2.assignments:
        if a.employee_id == "590373":
            a.account_name = "Fin_LKB9"
    rep = I.run_import(db, "WINSPEED", p2, dry_run=False)
    assert (rep["user_roles_added"], rep["user_roles_updated"], rep["user_roles_reactivated"],
            rep["user_roles_deactivated"]) == (0, 2, 1, 0)
    db.expire_all()
    assert {u.account_name for u in db.query(ArUserRole).filter(ArUserRole.employee_id == "590373")} == {"Fin_LKB9"}


def test_import_endpoint_defaults_to_dry_run(client, db, org):
    from tests.access_review.conftest import h

    org.user("IT", 3, 21)
    org.user("360002", 8, 1)
    wb = openpyxl.Workbook()
    wb.active.append(["รหัสพนักงาน", "Group Mail1"])
    wb.active.append([360002, "bod@menatransport.co.th"])
    buf = BytesIO()
    wb.save(buf)
    files = {"file": ("group-mails.xlsx", buf.getvalue(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post("/access-review/admin/import", headers=h("IT"), files=files, data={"system": "email"})
    assert r.status_code == 200, r.text
    assert r.json()["dry_run"] is True and r.json()["user_roles_added"] == 1
    assert db.query(ArUserRole).count() == 0
    assert client.post("/access-review/admin/import", headers=h("360002"), files=files,
                       data={"system": "EMAIL"}).status_code == 403
    bad = client.post("/access-review/admin/import", headers=h("IT"), files=files, data={"system": "SAP"})
    assert bad.status_code == 400
